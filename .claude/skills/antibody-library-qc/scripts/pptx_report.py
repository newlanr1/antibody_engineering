"""
Assemble a library QC PowerPoint deck from already-computed metrics and
already-rendered plot images. This module does not compute or interpret
anything — pass it real numbers and real conclusions, not placeholders.
"""

from pptx import Presentation
from pptx.util import Inches, Pt


def add_fitted_bullet_slide(prs, title_text, bullet_points):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title_shape = slide.shapes.title
    title_shape.text = title_text
    title_shape.text_frame.paragraphs[0].font.size = Pt(22)
    title_shape.text_frame.paragraphs[0].font.bold = True

    tf = slide.shapes.placeholders[1].text_frame
    tf.word_wrap = True
    tf.margin_left, tf.margin_right = Inches(0.5), Inches(0.5)
    tf.margin_top, tf.margin_bottom = Inches(0.2), Inches(0.2)

    tf.text = bullet_points[0]
    tf.paragraphs[0].font.size = Pt(14)
    tf.paragraphs[0].space_after = Pt(8)

    for pt in bullet_points[1:]:
        p = tf.add_paragraph()
        p.text = pt
        p.font.size = Pt(14)
        p.space_after = Pt(8)

    return slide


def add_image_slide(prs, title_text, img_path, width_in=9.0):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    title_shape = slide.shapes.title
    title_shape.text = title_text
    title_shape.text_frame.paragraphs[0].font.size = Pt(18)
    title_shape.text_frame.paragraphs[0].font.bold = True
    slide.shapes.add_picture(img_path, Inches(0.5), Inches(1.2), width=Inches(width_in))
    return slide


def _add_table(slide, headers, rows, left, top, width_in=9.0, col_widths_in=None,
                header_pt=12, body_pt=11, row_height_in=0.35):
    """Low-level table shape, shared by add_table_slide and add_summary_slide.
    `rows` is a list of row tuples/lists, each the same length as `headers`.
    Values are stringified as-is, so format numbers (commas, decimals) before
    passing them in."""
    n_rows = len(rows) + 1
    n_cols = len(headers)
    table_shape = slide.shapes.add_table(n_rows, n_cols, left, top,
                                          Inches(width_in), Inches(row_height_in * n_rows))
    table = table_shape.table

    if col_widths_in:
        for col_idx, w in enumerate(col_widths_in):
            table.columns[col_idx].width = Inches(w)

    for col_idx, header in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.text = str(header)
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(header_pt)
            p.font.bold = True

    for row_idx, row in enumerate(rows, start=1):
        for col_idx, value in enumerate(row):
            cell = table.cell(row_idx, col_idx)
            cell.text = str(value)
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(body_pt)

    return table_shape


def add_table_slide(prs, title_text, headers, rows, col_widths_in=None):
    """A title + data-table slide, with nothing else on it."""
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    title_shape = slide.shapes.title
    title_shape.text = title_text
    title_shape.text_frame.paragraphs[0].font.size = Pt(18)
    title_shape.text_frame.paragraphs[0].font.bold = True
    _add_table(slide, headers, rows, Inches(0.5), Inches(1.3), col_widths_in=col_widths_in)
    return slide


def _diversity_bullets(diversity_estimate):
    return [
        f"• Total Reads (N): {diversity_estimate['total_reads']:,}",
        f"• Unique Clones Observed (S_obs): {diversity_estimate['unique_clones_observed']:,}",
        f"• Singletons (f1) / Doubletons (f2): "
        f"{diversity_estimate['singletons_f1']:,} / {diversity_estimate['doubletons_f2']:,}",
        f"• Chao1 Estimate: {diversity_estimate['chao1_estimate']:,.0f} unique clones",
        f"• Poisson Estimate: {diversity_estimate['poisson_estimate']:,.0f} unique clones",
    ]


def add_summary_slide(prs, title_text, summary_bullets, diversity_estimate=None):
    """The library build summary slide. If `diversity_estimate` is given
    (diversity_estimation.py's diversity_estimate.json, already loaded), its
    point estimates are appended to the bullet list and its confidence-
    interval table is placed underneath — so every top-line number (QC
    efficiency *and* diversity) lives on one slide at the front of the deck
    instead of being scattered across separate slides later on.

    Uses the "Title Only" layout rather than a bullet placeholder so there's
    free canvas below the text to drop a table into.
    """
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    title_shape = slide.shapes.title
    title_shape.text = title_text
    title_shape.text_frame.paragraphs[0].font.size = Pt(22)
    title_shape.text_frame.paragraphs[0].font.bold = True

    bullets = list(summary_bullets)
    if diversity_estimate:
        bullets += _diversity_bullets(diversity_estimate)

    bullet_top_in = 1.1
    bullet_height_in = 0.26 * len(bullets) + 0.2
    txbox = slide.shapes.add_textbox(Inches(0.5), Inches(bullet_top_in),
                                      Inches(9.0), Inches(bullet_height_in))
    tf = txbox.text_frame
    tf.word_wrap = True
    tf.text = bullets[0]
    tf.paragraphs[0].font.size = Pt(13)
    tf.paragraphs[0].space_after = Pt(3)
    for b in bullets[1:]:
        p = tf.add_paragraph()
        p.text = b
        p.font.size = Pt(13)
        p.space_after = Pt(3)

    if diversity_estimate:
        headers = ["Confidence Level", "Chao1 Lower", "Chao1 Upper", "Poisson Lower", "Poisson Upper"]
        rows = [
            [level, f"{ci['chao1_lower']:,.0f}", f"{ci['chao1_upper']:,.0f}",
             f"{ci['poisson_lower']:,.0f}", f"{ci['poisson_upper']:,.0f}"]
            for level, ci in diversity_estimate["confidence_intervals"].items()
        ]
        table_top_in = bullet_top_in + bullet_height_in + 0.15
        _add_table(slide, headers, rows, Inches(0.5), Inches(table_top_in),
                   header_pt=10, body_pt=10, row_height_in=0.3)

    return slide


def build_qc_deck(output_path, title, subtitle, summary_metrics,
                   methodology_bullets, distribution_images,
                   pwm_images_by_gene, conclusion_bullets,
                   ml_bullets=None, ml_images=None, diversity_estimate=None,
                   developability_bullets=None, developability_images=None,
                   humanness_bullets=None, humanness_images=None,
                   structure_bullets=None, structure_table=None):
    """
    summary_metrics: dict of label -> value string, rendered as one bullet slide.
    methodology_bullets: list[str] describing how the QC was done.
    distribution_images: list of (slide_title, image_path) for global histogram slides.
    pwm_images_by_gene: list of (slide_title, image_path) for per-gene split PWM slides.
    conclusion_bullets: list[str] — findings actually observed in this run's
        results. Do not carry over conclusions from a different dataset/run.
    ml_bullets: optional list[str] summarizing KL-divergence / ESM-2 findings,
        rendered as its own slide before the image slides. Omit entirely if
        ml_bias_fitness.py wasn't run.
    ml_images: optional list of (slide_title, image_path) for ESM-2 perplexity
        scatter plots.
    diversity_estimate: optional dict loaded from diversity_estimation.py's
        diversity_estimate.json. Its point estimates and CI table are folded into
        the "Library Build Summary" slide (see add_summary_slide). Omit entirely
        if diversity_estimation.py wasn't run.
    developability_bullets: optional list[str] summarizing the sequence-based
        liability screen findings. developability_images: optional list of
        (slide_title, image_path), e.g. the liability-prevalence bar chart.
        Omit both if developability_screen.py wasn't run.
    humanness_bullets: optional list[str] summarizing the Sapiens humanness/
        immunogenicity findings. humanness_images: optional list of
        (slide_title, image_path), e.g. the humanness summary bar chart.
        Omit both if humanness_scoring.py wasn't run or skipped (e.g. the
        Sapiens environment wasn't set up).
    structure_bullets: optional list[str] about the structure-prediction shortlist.
        structure_table: optional (headers, rows) tuple from structure_prediction.py's
        manifest, rendered as a table slide (no 3D rendering -- see
        structure_prediction.py's module docstring for why). Omit both if that
        script wasn't run or skipped (e.g. the IgFold venv wasn't set up).
    """
    prs = Presentation()

    slide1 = prs.slides.add_slide(prs.slide_layouts[0])
    slide1.shapes.title.text = title
    slide1.placeholders[1].text = subtitle

    summary_pts = [f"• {k}: {v}" for k, v in summary_metrics.items()]
    add_summary_slide(prs, "Library Build Summary & Pairing Efficiency", summary_pts,
                       diversity_estimate=diversity_estimate)

    add_fitted_bullet_slide(prs, "QC Methodology", [f"• {b}" for b in methodology_bullets])

    for slide_title, img_path in distribution_images:
        add_image_slide(prs, slide_title, img_path)

    for slide_title, img_path in pwm_images_by_gene:
        add_image_slide(prs, slide_title, img_path)

    if ml_bullets:
        add_fitted_bullet_slide(prs, "Design Bias & ML Fitness Scoring", [f"• {b}" for b in ml_bullets])

    for slide_title, img_path in (ml_images or []):
        add_image_slide(prs, slide_title, img_path)

    if developability_bullets:
        add_fitted_bullet_slide(prs, "Developability & Liability Screening",
                                 [f"• {b}" for b in developability_bullets])

    for slide_title, img_path in (developability_images or []):
        add_image_slide(prs, slide_title, img_path)

    if humanness_bullets:
        add_fitted_bullet_slide(prs, "Humanness & Immunogenicity Risk (Sapiens)",
                                 [f"• {b}" for b in humanness_bullets])

    for slide_title, img_path in (humanness_images or []):
        add_image_slide(prs, slide_title, img_path)

    if structure_bullets:
        add_fitted_bullet_slide(prs, "Structure Prediction Shortlist",
                                 [f"• {b}" for b in structure_bullets])

    if structure_table:
        headers, rows = structure_table
        add_table_slide(prs, "Structure Prediction Shortlist", headers, rows)

    if conclusion_bullets:
        add_fitted_bullet_slide(prs, "Conclusions & Recommendations",
                                 [f"• {b}" for b in conclusion_bullets])

    prs.save(output_path)
    return output_path
